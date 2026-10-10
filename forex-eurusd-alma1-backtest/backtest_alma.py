import ctypes
import csv
import json
import math
import os
import subprocess
from dataclasses import asdict, dataclass
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font


ROOT = Path(__file__).resolve().parent
WORK = ROOT / 'work'
WORK.mkdir(exist_ok=True)
ATTACH = ROOT / 'attachments'
CSV_PATH = ATTACH / 'ALMA - P ONLY.csv'
TEMPLATE = ATTACH / 'RÉSULTATS INDICATEURS.xlsx'
PIP = 0.0001
START_CAPITAL = 100_000.0
MAX_RISK_USD = 1_000.0
LEVERAGE = 5.0
LOT_STEP = 0.01
PIP_USD_PER_LOT = 10.0
ROUND_COST_PIPS = 0.8 + 7.0 / PIP_USD_PER_LOT
BLOCK_MIN = [200, 70, 70, 70]


@dataclass(frozen=True)
class Policy:
    family: str
    stop_kind: int
    stop_value: float
    rr: float
    max_bars: int

    @property
    def name(self):
        sl = f'{self.stop_value:g} pips' if self.stop_kind == 0 else 'ATR14 (n−1)'
        cap = 'sans limite' if self.max_bars == 0 else f'{self.max_bars} bougies'
        return f'SL {sl}; RR 1:{self.rr:g}; durée {cap}'


def make_policies():
    policies = []
    times = [0, 12, 24, 36]
    for stop in [10, 20, 30]:
        for rr in [2, 3, 4, 5, 8, 10]:
            for cap in times:
                policies.append(Policy('PIPS', 0, float(stop), float(rr), cap))
    for rr in [2, 3, 4, 5]:
        for cap in times:
            policies.append(Policy('PIPS', 0, 40.0, float(rr), cap))
    for rr in [2, 3, 4, 5, 8, 10]:
        for cap in times:
            policies.append(Policy('ATR14', 1, 1.0, float(rr), cap))
    return policies


def load_prices():
    df = pd.read_csv(CSV_PATH, sep=';', decimal=',')
    cols = ['OPEN', 'HIGH', 'LOW', 'CLOSE', 'ATR14 ', 'ALMA-1']
    for c in cols:
        df[c] = pd.to_numeric(df[c], errors='coerce')
    complete = df[['OPEN', 'HIGH', 'LOW', 'CLOSE']].notna().all(axis=1).to_numpy()
    n = int(complete.sum())
    # Keep the complete OHLC history only. The source ends with a partial row
    # containing an indicator value but no executable candle.
    df = df.iloc[:n].copy()
    o = np.ascontiguousarray(df['OPEN'].to_numpy(dtype=np.float64))
    h = np.ascontiguousarray(df['HIGH'].to_numpy(dtype=np.float64))
    l = np.ascontiguousarray(df['LOW'].to_numpy(dtype=np.float64))
    c = np.ascontiguousarray(df['CLOSE'].to_numpy(dtype=np.float64))
    atr = df['ATR14 '].to_numpy(dtype=np.float64)
    atr_prev = np.full(n, np.nan, dtype=np.float64)
    atr_prev[1:] = atr[:-1]
    atr_prev = np.ascontiguousarray(atr_prev)
    indicator = np.ascontiguousarray(df['ALMA-1'].to_numpy(dtype=np.float64))
    means = np.full(n, np.nan, dtype=np.float64)
    means[2:] = (indicator[2:] + indicator[1:-1] + indicator[:-2]) / 3.0
    raw = indicator[:n // 2]
    raw = raw[np.isfinite(raw)]
    return df, n, o, h, l, c, atr_prev, indicator, means, raw


def compile_outcomes(n, o, h, l, c, atr_prev, policies):
    src = ROOT / 'outcomes.cpp'
    libpath = WORK / 'outcomes.so'
    subprocess.run(['g++', '-O3', '-std=c++17', '-fPIC', '-shared', str(src), '-o', str(libpath)], check=True)
    lib = ctypes.CDLL(str(libpath))
    fn = lib.precompute_outcomes
    D = ctypes.POINTER(ctypes.c_double)
    I = ctypes.POINTER(ctypes.c_int)
    I32 = ctypes.POINTER(ctypes.c_int32)
    fn.argtypes = [ctypes.c_int, D, D, D, D, D, ctypes.c_int, I, D, D, I, I32, D]
    fn.restype = None
    stop_kind = np.ascontiguousarray([p.stop_kind for p in policies], dtype=np.int32)
    stop_value = np.ascontiguousarray([p.stop_value for p in policies], dtype=np.float64)
    rr = np.ascontiguousarray([p.rr for p in policies], dtype=np.float64)
    max_bars = np.ascontiguousarray([p.max_bars for p in policies], dtype=np.int32)
    size = len(policies) * 2 * n
    exit_index = np.empty(size, dtype=np.int32)
    gross_r = np.empty(size, dtype=np.float64)
    fn(n, o.ctypes.data_as(D), h.ctypes.data_as(D), l.ctypes.data_as(D), c.ctypes.data_as(D),
       atr_prev.ctypes.data_as(D), len(policies), stop_kind.ctypes.data_as(I),
       stop_value.ctypes.data_as(D), rr.ctypes.data_as(D), max_bars.ctypes.data_as(I),
       exit_index.ctypes.data_as(I32), gross_r.ctypes.data_as(D))
    return exit_index.reshape((len(policies), 2, n)), gross_r.reshape((len(policies), 2, n))


def stop_pips(policy, entry, atr_prev):
    return policy.stop_value if policy.stop_kind == 0 else atr_prev[entry] / PIP


def r_squared(net_r):
    m = len(net_r)
    if m < 2:
        return 0.0
    y = np.cumsum(net_r, dtype=np.float64)
    x = np.arange(1, m + 1, dtype=np.float64)
    dx = x - x.mean()
    dy = y - y.mean()
    denom = float(np.dot(dx, dx) * np.dot(dy, dy))
    if denom <= 0:
        return 1.0 if np.allclose(y, y[0]) else 0.0
    corr = float(np.dot(dx, dy)) / math.sqrt(denom)
    return corr * corr


def run_strategy(events, policy_i, direction, policy, exits, gross, o, atr_prev, n):
    next_entry = 0
    trades_gross = []
    trades_net = []
    blocks_gross = [[] for _ in range(4)]
    blocks_net = [[] for _ in range(4)]
    equity = START_CAPITAL
    peak = equity
    max_dd = 0.0
    for entry in events:
        if entry < next_entry:
            continue
        ix = int(exits[policy_i, direction, entry])
        if ix < entry:
            continue
        sp = stop_pips(policy, entry, atr_prev)
        if not (sp > 0.0 and math.isfinite(sp)):
            continue
        risk_lots = MAX_RISK_USD / (sp * PIP_USD_PER_LOT)
        margin_lots = equity * LEVERAGE / (100_000.0 * o[entry])
        lots = min(risk_lots, margin_lots)
        lots = math.floor((lots + 1e-10) / LOT_STEP) * LOT_STEP
        if lots < LOT_STEP:
            continue
        gr = float(gross[policy_i, direction, entry])
        nr = gr - ROUND_COST_PIPS / sp
        risk_usd = lots * sp * PIP_USD_PER_LOT
        pnl = nr * risk_usd
        equity += pnl
        peak = max(peak, equity)
        dd = (peak - equity) / peak if peak > 0 else 1.0
        max_dd = max(max_dd, dd)
        trades_gross.append(gr)
        trades_net.append(nr)
        block = 0 if entry < int(0.4*n) else 1 if entry < int(0.6*n) else 2 if entry < int(0.8*n) else 3
        blocks_gross[block].append(gr)
        blocks_net[block].append(nr)
        next_entry = ix + 1
    count = len(trades_net)
    gross_mean = float(np.mean(trades_gross)) if count else float('nan')
    net_mean = float(np.mean(trades_net)) if count else float('nan')
    block_counts = [len(x) for x in blocks_net]
    block_means = [float(np.mean(x)) if x else float('nan') for x in blocks_net]
    return {
        'trades': count,
        'gross_mean_r': gross_mean,
        'net_mean_r': net_mean,
        'max_drawdown_pct': 100.0 * max_dd,
        'r2': r_squared(np.asarray(trades_net, dtype=np.float64)),
        'block_counts': block_counts,
        'block_net_mean_r': block_means,
        'ending_equity': equity,
        'policy': policy.name,
        'direction': 'LONG ONLY' if direction == 0 else 'SHORT ONLY',
        'policy_index': int(policy_i),
        'direction_index': int(direction),
    }


def evaluate_candidate(cond, lo_pct, hi_pct, thresholds, means, policies, exits, gross, o, atr_prev, n):
    if cond == 'LT':
        threshold = float(thresholds[lo_pct])
        event_mask = np.isfinite(means) & (means < threshold)
        description = f'< Q{lo_pct} ({threshold:.8f})'
    elif cond == 'GT':
        threshold = float(thresholds[lo_pct])
        event_mask = np.isfinite(means) & (means > threshold)
        description = f'> Q{lo_pct} ({threshold:.8f})'
    else:
        lower = float(thresholds[lo_pct])
        upper = float(thresholds[hi_pct])
        event_mask = np.isfinite(means) & (means > lower) & (means < upper)
        description = f'Q{lo_pct}–Q{hi_pct} ({lower:.8f}–{upper:.8f})'
    events = np.flatnonzero(event_mask).astype(np.int32)
    if not len(events):
        return {'cond': cond, 'lo_pct': lo_pct, 'hi_pct': hi_pct, 'description': description,
                'best': None, 'tried_configs': 0, 'events': 0}
    best = None
    valid_best = None
    tried = 0
    for direction in (0, 1):
        for policy_i, policy in enumerate(policies):
            result = run_strategy(events, policy_i, direction, policy, exits, gross, o, atr_prev, n)
            tried += 1
            result.update({'cond': cond, 'lo_pct': int(lo_pct), 'hi_pct': None if hi_pct is None else int(hi_pct),
                           'entry_threshold': description, 'events': int(len(events))})
            if result['trades'] >= 1000 and (best is None or result['net_mean_r'] > best['net_mean_r']):
                best = result
            if result['trades'] >= 1000:
                strict_ok, _ = check_validation(result, tolerance=False)
                tolerant_ok, _ = check_validation(result, tolerance=True)
                if (strict_ok or tolerant_ok) and (valid_best is None or result['net_mean_r'] > valid_best['net_mean_r']):
                    valid_best = result
    return {'cond': cond, 'lo_pct': int(lo_pct), 'hi_pct': None if hi_pct is None else int(hi_pct),
            'description': description, 'best': best, 'valid_best': valid_best,
            'tried_configs': tried, 'events': int(len(events))}


def check_validation(best, tolerance=False):
    if not best or best['trades'] < 1000:
        return False, 'moins de 1 000 groupes indépendants sur l’ensemble'
    for i, (cnt, mean) in enumerate(zip(best['block_counts'], best['block_net_mean_r'])):
        if cnt < BLOCK_MIN[i]:
            return False, f'bloc {i+1}: {cnt} groupes (< {BLOCK_MIN[i]})'
        floor = 0.375 if tolerance and i > 0 else 0.5
        if not math.isfinite(mean) or mean < floor:
            return False, f'bloc {i+1}: R net moyen {mean:.4f} < {floor:.3f}'
    if best['net_mean_r'] < 0.5:
        return False, f'R net moyen global {best["net_mean_r"]:.4f} < 0.500'
    if best['max_drawdown_pct'] >= 20.0:
        return False, f'drawdown {best["max_drawdown_pct"]:.2f}% ≥ 20%'
    if best['r2'] < 0.77:
        return False, f'R² {best["r2"]:.4f} < 0.77'
    return True, 'validé'


def candidate_key(c):
    return (c['cond'], c['lo_pct'], c['hi_pct'])


def coarse_candidates(qs):
    out = []
    for cond in ['LT', 'GT']:
        for p in range(10, 100, 10):
            out.append((cond, p, None))
    for lo in range(10, 100, 10):
        for hi in range(lo+10, 100, 10):
            out.append(('BETWEEN', lo, hi))
    return out


def refinement_candidates(summaries):
    # Refine locally around the best coarse zones found for each direction and condition.
    seeds = {}
    for cond in ['LT', 'GT', 'BETWEEN']:
        allc = [c for c in summaries if c['cond'] == cond and c.get('best') is not None]
        for direction in ['LONG ONLY', 'SHORT ONLY']:
            # The candidate summary retains its best direction. Keep the top three
            # qualifying coarse zones for each entry condition.
            cand = sorted([c for c in allc if c['best']['direction'] == direction],
                          key=lambda c: c['best']['net_mean_r'], reverse=True)[:3]
            for c in cand:
                seeds.setdefault(cond, set()).add((c['lo_pct'], c['hi_pct']))
    out = set()
    for cond, pairset in seeds.items():
        for lo, hi in pairset:
            lows = range(max(1, lo-5), min(99, lo+5)+1)
            if cond in ('LT', 'GT'):
                for p in lows:
                    out.add((cond, p, None))
            else:
                highs = range(max(2, hi-5), min(99, hi+5)+1)
                for lp in lows:
                    for hp in highs:
                        if lp < hp:
                            out.add((cond, lp, hp))
    # Preserve the requested search order: all LT, then GT, then BETWEEN.
    return sorted(out, key=lambda x: ({'LT':0,'GT':1,'BETWEEN':2}[x[0]], x[1], x[2] or 0))


def choose_best_three(summaries):
    out = {'LT': [], 'GT': [], 'BETWEEN': []}
    by_key = {}
    for summary in summaries:
        best = summary.get('best')
        if best is None:
            continue
        key = candidate_key(summary)
        if key not in by_key or best['net_mean_r'] > by_key[key]['net_mean_r']:
            by_key[key] = best
    for cond in out:
        ranked = sorted([x for x in by_key.items() if x[0][0] == cond],
                        key=lambda x: x[1]['net_mean_r'], reverse=True)
        out[cond] = [v for _, v in ranked[:3]]
    return out


def save_outputs(summaries, df, policies, n, raw, qs):
    # Store generated workbooks and search logs in the local work directory.
    output = WORK / 'deliverables'
    output.mkdir(parents=True, exist_ok=True)
    results_path = output / 'RÉSULTATS INDICATEURS.xlsx'
    reject_path = output / 'historique_rejets_ALMA-1.csv'
    best_by = choose_best_three(summaries)
    all_tests = []
    validated = []
    for summary in summaries:
        best = summary.get('best')
        chosen = summary.get('valid_best') or best
        if chosen is None:
            status, reason = False, 'moins de 1 000 groupes indépendants sur l’ensemble'
        else:
            status, reason = check_validation(chosen, tolerance=False)
            tol_status, tol_reason = check_validation(chosen, tolerance=True)
            if tol_status and not status:
                status, reason = True, 'validation avec tolérance: ' + tol_reason
            if status:
                validated.append(chosen)
        row = {
            'condition': summary['cond'], 'percentile_bas': summary['lo_pct'],
            'percentile_haut': summary['hi_pct'], 'seuil': summary['description'],
            'effectif_indicateur': summary['events'], 'configurations_sortie_testées': summary['tried_configs'],
            'meilleur_sens': chosen['direction'] if chosen else '',
            'meilleure_sortie': chosen['policy'] if chosen else '',
            'trades_indépendants': chosen['trades'] if chosen else 0,
            'R_brut_moyen': chosen['gross_mean_r'] if chosen else '',
            'R_net_moyen': chosen['net_mean_r'] if chosen else '',
            'drawdown_max_pct': chosen['max_drawdown_pct'] if chosen else '',
            'R2': chosen['r2'] if chosen else '',
            'groupes_blocs': json.dumps(chosen['block_counts'], ensure_ascii=False) if chosen else '',
            'R_net_blocs': json.dumps(chosen['block_net_mean_r'], ensure_ascii=False) if chosen else '',
            'statut_validation': 'VALIDÉ' if status else 'REJETÉ',
            'motif': reason,
        }
        all_tests.append(row)
    with open(reject_path, 'w', newline='', encoding='utf-8-sig') as f:
        fields = list(all_tests[0].keys()) if all_tests else ['condition','seuil','statut_validation']
        writer = csv.DictWriter(f, fieldnames=fields, delimiter=';')
        writer.writeheader()
        writer.writerows([r for r in all_tests if r['statut_validation'] != 'VALIDÉ'])

    wb = load_workbook(TEMPLATE)
    ws = wb['RÉSULTATS INDICATEURS']
    ws.cell(row=1, column=2, value='ALMA-1')
    rows_by_cond = {'LT': [(3,4,5,6,7),(9,10,11,12,13),(15,16,17,18,19)],
                    'GT': [(21,22,23,24,25),(27,28,29,30,31),(33,34,35,36,37)],
                    'BETWEEN': [(39,40,41,42,43),(45,46,47,48,49),(51,52,53,54,55)]}
    for cond, records in best_by.items():
        for rank, record in enumerate(records):
            rankrow, thresholdrow, grossrow, ddrow, stratrow = rows_by_cond[cond][rank]
            ws.cell(rankrow, 2, rank+1)
            ws.cell(thresholdrow, 2, record['entry_threshold'])
            ws.cell(grossrow, 2, record['gross_mean_r'])
            ws.cell(ddrow, 2, record['max_drawdown_pct']/100.0)
            ws.cell(ddrow, 2).number_format = '0.00%'
            ws.cell(stratrow, 2, f"{record['direction']}; {record['policy']}")
    ws.column_dimensions['A'].width = 36
    ws.column_dimensions['B'].width = 43
    for row in ws.iter_rows(min_row=1, max_col=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    if validated:
        details = wb.create_sheet('Stratégie validée')
        details.append(['Seuil','Sens','Sortie','Trades indépendants','R brut moyen','R net moyen','Drawdown max','R²','Trades par bloc','R net moyen par bloc'])
        for v in validated:
            details.append([v['entry_threshold'],v['direction'],v['policy'],v['trades'],v['gross_mean_r'],
                            v['net_mean_r'],v['max_drawdown_pct']/100,v['r2'],json.dumps(v['block_counts']),
                            json.dumps(v['block_net_mean_r'])])
        details.freeze_panes='A2'
        details.auto_filter.ref=details.dimensions
        for cell in details[1]:
            cell.font=Font(bold=True); cell.alignment=Alignment(wrap_text=True,vertical='top')
        for col,width in enumerate([36,15,44,18,14,14,14,12,22,30],1):
            details.column_dimensions[chr(64+col)].width=width
        for col in ['E','F','H']:
            for cell in details[col][1:]:
                if isinstance(cell.value,(int,float)): cell.number_format='0.0000'
        for cell in details['G'][1:]:
            if isinstance(cell.value,(int,float)): cell.number_format='0.00%'
    wb.save(results_path)
    return str(results_path), str(reject_path), validated, best_by


def main():
    df, n, o, h, l, c, atr_prev, indicator, means, raw = load_prices()
    policies = make_policies()
    print(f'Données: {n} bougies complètes; {len(policies)} variantes de sortie; {len(raw)} valeurs ALMA-1 valides dans la première moitié.', flush=True)
    print(f'ALMA-1 min/max (tout l’historique): {np.nanmin(indicator):.8f} / {np.nanmax(indicator):.8f}', flush=True)
    print('Distribution globale de ALMA-1:', flush=True)
    for p in [10,20,30,40,50,60,70,80,90]:
        print(f'Q{p}: {np.nanpercentile(indicator,p):.8f}', flush=True)
    print('Déciles sur les valeurs brutes valides de la première moitié:', flush=True)
    for p in range(0,101,10):
        print(f'Q{p}: {np.percentile(raw,p):.8f}', flush=True)
    qv = np.percentile(raw, np.arange(1,100,dtype=float))
    thresholds = {p: float(qv[p-1]) for p in range(1,100)}
    exits, gross = compile_outcomes(n,o,h,l,c,atr_prev,policies)
    print('Simulation OHLC des variantes de sortie pré-calculée.', flush=True)
    summaries = []
    tested = set()
    checkpoint = WORK / 'checkpoint.jsonl'
    if checkpoint.exists():
        checkpoint.unlink()

    def test_batch(candidates, label):
        ordered = sorted(set(candidates), key=lambda x: ({'LT':0,'GT':1,'BETWEEN':2}[x[0]],x[1],x[2] or 0))
        print(f'{label}: {len(ordered)} combinaisons de seuil.', flush=True)
        for i,(cond,lo,hi) in enumerate(ordered,1):
            key=(cond,lo,hi)
            if key in tested: continue
            s=evaluate_candidate(cond,lo,hi,thresholds,means,policies,exits,gross,o,atr_prev,n)
            summaries.append(s); tested.add(key)
            with open(checkpoint,'a',encoding='utf-8') as f:
                f.write(json.dumps(s,ensure_ascii=False,allow_nan=True)+'\n')
            if i%25==0 or i==len(ordered):
                print(f'  {label}: {i}/{len(ordered)}',flush=True)
            valid=s.get('valid_best')
            if valid:
                print('CANDIDAT VALIDÉ',json.dumps(valid,ensure_ascii=False),flush=True)
                return valid
        return None

    found=test_batch(coarse_candidates(thresholds),'Balayage déciles')
    if not found:
        fine=refinement_candidates(summaries)
        found=test_batch(fine,'Affinage local à 1 percentile')
    results_path,reject_path,validated,best_by=save_outputs(summaries,df,policies,n,raw,thresholds)
    state={'n_bars':n,'raw_valid_first_half':len(raw),'tested_thresholds':len(tested),
           'strategies_validated':len(validated),'best_by_condition':best_by,
           'results_path':results_path,'reject_path':reject_path}
    with open(WORK/'progress.json','w',encoding='utf-8') as f: json.dump(state,f,ensure_ascii=False,indent=2)
    print('TERMINÉ',json.dumps(state,ensure_ascii=False),flush=True)


if __name__ == '__main__':
    main()
